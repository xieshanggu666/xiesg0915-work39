"""多专业协同核查端到端测试（不依赖 pytest，可直接运行）。

用法::

    python tests/test_coordination.py
"""

from __future__ import annotations

import json
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tools.make_sample_coordination import make_coordination_sample  # noqa: E402
from ifc_audit.coordination import (  # noqa: E402
    run_coordination, resolve_settings, resolve_coord_gate,
    assign_issue, fix_issue, verify_issue,
    reject_issue, CoordWorkflowError, parse_owner_items,
    apply_sla_sweep, backfill_ledger_sla,
)
from ifc_audit.coordination_model import (  # noqa: E402
    CoordinationLedger, CoordIssue, DISC_ARCH, DISC_STRUCT, DISC_MEP,
    KIND_HARD_CLASH, KIND_OPENING_MISSING, KIND_OPENING_MISMATCH,
    KIND_OPENING_UNUSED, STATUS_OPEN, STATUS_FIXED, STATUS_VERIFIED,
    STATUS_REJECTED, discipline_from_filename,
)
from datetime import datetime, timedelta  # noqa: E402
from ifc_audit.batch import run_batch_with_config  # noqa: E402
from ifc_audit import coordination_report  # noqa: E402


def _kind_counts(result, active_only=True):
    out = {}
    for i in result.issues:
        if active_only and not i.active:
            continue
        out[i.kind] = out.get(i.kind, 0) + 1
    return out


def run() -> int:
    failures = []

    def check(cond, msg):
        print(("PASS " if cond else "FAIL ") + msg)
        if not cond:
            failures.append(msg)

    with tempfile.TemporaryDirectory() as td:
        sample = os.path.join(td, "sample")
        paths = make_coordination_sample(sample)
        owners = {DISC_STRUCT: "张结构", DISC_MEP: "李机电",
                  DISC_ARCH: "王建筑"}
        ledger_path = os.path.join(td, "history", "测试项目",
                                   "coordination_ledger.json")

        # ---- 1) 四类问题检测 ----
        result = run_coordination(
            [sample], project="测试项目", label="v1",
            owners=owners, ledger_path=ledger_path)
        check(len(result.files) == 3, "纳入建筑/结构/机电 3 份专业模型")
        discs = {f.discipline for f in result.files}
        check(discs == {DISC_ARCH, DISC_STRUCT, DISC_MEP},
              f"文件名正确识别专业（实际 {discs}）")
        counts = _kind_counts(result)
        check(counts.get(KIND_HARD_CLASH) == 1,
              f"检出 1 处硬碰撞（实际 {counts.get(KIND_HARD_CLASH, 0)}）")
        check(counts.get(KIND_OPENING_MISSING) == 1,
              f"检出 1 处洞口缺失（实际 {counts.get(KIND_OPENING_MISSING, 0)}）")
        check(counts.get(KIND_OPENING_MISMATCH) == 1,
              f"检出 1 处洞口规格不符（实际 {counts.get(KIND_OPENING_MISMATCH, 0)}）")
        check(counts.get(KIND_OPENING_UNUSED) == 1,
              f"检出 1 处洞口未使用（实际 {counts.get(KIND_OPENING_UNUSED, 0)}）")

        # ---- 2) 责任分派 ----
        clash = next(i for i in result.issues if i.kind == KIND_HARD_CLASH)
        check(clash.owner_discipline == DISC_MEP and clash.owner == "李机电",
              f"硬碰撞派机电/李机电（实际 {clash.owner_discipline}/{clash.owner}）")
        missing = next(i for i in result.issues
                       if i.kind == KIND_OPENING_MISSING)
        check(missing.owner_discipline == DISC_ARCH,
              "墙洞缺失派建筑专业")
        unused = next(i for i in result.issues
                      if i.kind == KIND_OPENING_UNUSED)
        check(unused.owner_discipline == DISC_MEP,
              "洞口未使用派机电专业核对路由")
        for i in result.issues:
            check(i.status == STATUS_OPEN, f"{i.issue_id} 初始状态待整改")

        # ---- 3) 默认门禁阻断（错误类零容忍）----
        check(not result.gate_passed, "默认协同门禁：错误类未清零 -> 阻断")
        failed = [r for r in result.gate_rules if not r["passed"]]
        check(any(r["key"] == "coord_max_clash_active" for r in failed),
              "失败规则为未闭环碰撞/缺洞数")

        # 宽松门禁允许 3 项错误类 -> 通过
        loose = run_coordination(
            [sample], project="测试项目", owners=owners,
            ledger_path=ledger_path, gate_profile="loose")
        check(loose.gate_passed, "宽松协同门禁下放行")

        # ---- 4) 工单流转：派单 -> 整改 -> 复核 ----
        ledger = CoordinationLedger.load(ledger_path)
        cid = clash.issue_id
        assign_issue(ledger, cid, "赵结构", discipline=DISC_STRUCT,
                     by="协调人", note="碰撞改由结构调整梁高")
        issue = next(i for i in ledger.issues.values() if i.issue_id == cid)
        check(issue.owner_discipline == DISC_STRUCT and issue.owner == "赵结构",
              "改派到结构/赵结构")
        fix_issue(ledger, cid, by="赵结构", note="梁高调整完成")
        check(ledger.get(issue.fingerprint).status == STATUS_FIXED,
              "整改后进入待复核")
        # 待整改工单不能直接驳回
        mid = missing.issue_id
        try:
            reject_issue(ledger, mid, by="王建筑", note="没补好")
            check(False, "待整改状态 reject 应报错")
        except CoordWorkflowError:
            check(True, "待整改状态 reject 被拒绝（状态机保护）")
        verify_issue(ledger, cid, by="李机电", note="复核通过")
        check(ledger.get(issue.fingerprint).status == STATUS_VERIFIED,
              "复核通过后闭环")
        check(ledger.get(issue.fingerprint).verified_by == "李机电",
              "记录复核人")
        ledger.save(ledger_path)

        # ---- 11) 整改时限：派单即设时限，未到期不升级 ----
        sla_dir = os.path.join(td, "sla")
        make_coordination_sample(sla_dir)
        sla_ledger = os.path.join(td, "history", "时限项目",
                                  "coordination_ledger.json")
        sla_res = run_coordination(
            [sla_dir], project="时限项目", owners=owners,
            ledger_path=sla_ledger,
            gate_overrides={"coord_fix_sla_hours": 48})
        for i in sla_res.issues:
            check(i.sla_hours == 48 and bool(i.due_at),
                  f"{i.issue_id} 派单即设 48h 整改时限")
            check(not i.is_overdue() and not i.escalated,
                  f"{i.issue_id} 新建工单未超期、未升级")
        check(sla_res.summary()["issues_overdue"] == 0,
              "未到期：超期数为 0")

        # 手工把截止时间拨到 10 小时前 -> 扫描后自动升级到 L1
        sla_ld = CoordinationLedger.load(sla_ledger)
        target = next(i for i in sla_ld.issues.values() if i.active)
        target.due_at = (datetime.now() - timedelta(hours=10)
                         ).isoformat(timespec="seconds")
        now = datetime.now()
        escalated = apply_sla_sweep(sla_ld, 48, "B-TEST", now)
        check(len(escalated) == 1 and escalated[0].issue_id == target.issue_id,
              "超期未整改 -> 自动升级（1 张）")
        check(target.escalated and target.escalation_level == 1,
              "首次超期升级到 L1 责任专业负责人")
        check(any(h["action"] == "escalate" for h in target.history),
              "自动升级写入流转记录")
        # 幂等：再次扫描（仍只超 10h）不重复升级
        again = apply_sla_sweep(sla_ld, 48, "B-TEST",
                                now + timedelta(hours=1))
        check(again == [] and target.escalation_level == 1,
              "时限扫描幂等：不重复升级")
        # 再越过一个时限周期 -> 升级到 L2
        lvl2 = apply_sla_sweep(
            sla_ld, 48, "B-TEST", now + timedelta(hours=60))
        check(target in lvl2 and target.escalation_level == 2,
              "超期超过一个时限周期 -> 升级 L2 项目协调")
        sla_ld.save(sla_ledger)

        # 超期工单进入协同放行判定（strict：超期零容忍）
        sla_res2 = run_coordination(
            [sla_dir], project="时限项目", owners=owners,
            ledger_path=sla_ledger, gate_profile="strict")
        od_rules = [r for r in sla_res2.gate_rules
                    if r["key"] == "coord_max_overdue_active"]
        check(bool(od_rules) and not od_rules[0]["passed"],
              "超期工单计入协同放行判定（strict 门禁阻断）")
        check(not sla_res2.gate_passed, "超期零容忍 -> 协同门禁不通过")
        check(sla_res2.summary()["issues_overdue"] >= 1,
              "汇总统计超期工单数")

        # ---- 12) 驳回重排时限、升级清零 ----
        other = next(i for i in sla_ld.issues.values()
                     if i.issue_id != target.issue_id and i.active)
        oid = other.issue_id
        fix_issue(sla_ld, oid, by="王建筑", note="补洞完成")
        reject_issue(sla_ld, oid, by="李机电", note="洞口偏", sla_hours=48)
        o = sla_ld.get(other.fingerprint)
        check(o.status == STATUS_REJECTED and o.escalation_level == 0
              and not o.escalated,
              "驳回后清零升级标记")
        from ifc_audit.coordination_model import parse_dt
        due = parse_dt(o.due_at)
        check(abs((due - datetime.now()).total_seconds() / 3600.0 - 48) < 0.05,
              "驳回后按新一轮整改重排 48h 时限")

        # ---- 13) 历史台账兼容：旧工单缺时限，合并时补录且不追溯 ----
        hist_dir = os.path.join(td, "legacy")
        make_coordination_sample(hist_dir)
        legacy_ledger = os.path.join(td, "history", "历史项目",
                                     "coordination_ledger.json")
        r0 = run_coordination(
            [hist_dir], project="历史项目", owners=owners,
            ledger_path=legacy_ledger,
            gate_overrides={"coord_fix_sla_hours": 0})  # 模拟旧版无时限
        old_issue = next(iter(r0.issues))
        check(not old_issue.due_at and old_issue.sla_hours == 0,
              "旧版工单没有整改时限字段")
        # 把台账伪装成 30 天前创建的历史遗留工单
        hist_ld = CoordinationLedger.load(legacy_ledger)
        ancient = (datetime.now() - timedelta(days=30)
                   ).isoformat(timespec="seconds")
        for i in hist_ld.issues.values():
            i.created_at = ancient
            i.updated_at = ancient
            i.due_at = ""
            i.sla_hours = 0.0
        n = backfill_ledger_sla(hist_ld, 72, "B-LEGACY")
        check(n == len(hist_ld.issues),
              f"历史活动工单全部补录时限（{n} 张）")
        for i in hist_ld.issues.values():
            check(not i.is_overdue(),
                  "补录自当前时刻起算：历史工单不被立即判超期")
            check(any(h["action"] == "sla_backfill" for h in i.history),
                  "时限补录写入流转记录")
        # 已闭环工单不补录
        hist_closed = CoordinationLedger(project="历史项目")
        closed = CoordIssue(issue_id="COORD-9001", fingerprint="xclosed",
                            status=STATUS_VERIFIED)
        hist_closed.issues[closed.fingerprint] = closed
        check(backfill_ledger_sla(hist_closed, 72) == 0,
              "已闭环历史工单不补时限")

        # 旧台账（缺新字段）可正常加载并参与新一轮核查
        hist_ld.save(legacy_ledger)
        r1 = run_coordination(
            [hist_dir], project="历史项目", owners=owners,
            ledger_path=legacy_ledger)
        check(all(i.sla_hours == 72 for i in r1.issues if i.active),
              "旧台账加载后新一轮核查自动带上默认 72h 时限")

        # ---- 14) 报告含时限/超期/升级列 ----
        out_sla = os.path.join(td, "out_sla")
        exported_sla = coordination_report.export_all(sla_res2, out_sla)
        from openpyxl import load_workbook
        wb2 = load_workbook(exported_sla["excel"])
        ws = wb2["碰撞与洞口工单"]
        head = [c.value for c in ws[1]]
        check({"整改时限(h)", "整改截止", "剩余/超期(h)", "升级状态"}
              <= set(head),
              "协同 Excel 工单表含时限/截止/超期/升级列")
        import csv as _csv
        with open(exported_sla["issues_csv"], encoding="utf-8-sig") as f:
            csv_head = next(_csv.reader(f))
        check({"整改时限h", "整改截止", "剩余或超期h", "升级级别"}
              <= set(csv_head),
              "工单 CSV 含时限列")
        with open(exported_sla["arch_writeback"], encoding="utf-8") as f:
            wj = json.load(f)
        check("issues_overdue" in wj and "issues_escalated" in wj
              and "fix_sla_hours" in wj,
              "建筑侧回写结论含超期/升级/时限字段")


        # ---- 5) 复核驳回流转 ----
        ledger = CoordinationLedger.load(ledger_path)
        fix_issue(ledger, mid, by="王建筑", note="已补洞口")
        reject_issue(ledger, mid, by="李机电", note="洞口位置仍偏")
        check(ledger.get(missing.fingerprint).status == STATUS_REJECTED,
              "待复核可驳回，退回整改")
        fix_issue(ledger, mid, by="王建筑", note="重新定位补洞")
        check(ledger.get(missing.fingerprint).status == STATUS_FIXED,
              "驳回后可再次报整改")
        ledger.save(ledger_path)

        # ---- 6) 报告导出不报错且含工单 ----
        out = os.path.join(td, "out")
        exported = coordination_report.export_all(result, out)
        check(all(os.path.exists(p) for p in exported.values()),
              f"协同 Excel/CSV/JSON/回写均导出（{len(exported)} 个文件）")
        from openpyxl import load_workbook
        wb = load_workbook(exported["excel"])
        check({"协同概览", "碰撞与洞口工单", "专业模型清单",
               "判定参数与门禁"} <= set(wb.sheetnames),
              f"协同 Excel 含 4 张表（实际 {wb.sheetnames}）")
        with open(exported["arch_writeback"], encoding="utf-8") as f:
            wb_json = json.load(f)
        check(wb_json["issues_active"] == 4, "建筑侧回写结论含未闭环数")
        per_unit = {r["unit"]: r for r in wb_json["per_arch_unit"]}
        check("样例-建筑模型" in per_unit, "回写按建筑单体拆分")

        # ---- 7) 重新核查：问题消失自动闭环、复核通过后回归自动重开 ----
        # 7a) 在“仍有碰撞、但缺洞已补”的模型上：已复核的碰撞仍检出 -> 回归重开
        partial_dir = os.path.join(td, "partial")
        make_partial_sample(partial_dir, paths)
        run_coordination(
            [partial_dir], project="测试项目", owners=owners,
            ledger_path=ledger_path)
        ledger2 = CoordinationLedger.load(ledger_path)
        clash2 = ledger2.get(clash.fingerprint)
        check(clash2.status == STATUS_OPEN,
              f"复核通过后仍检出 -> 回归重开（实际 {clash2.status}）")
        check(any(h["action"] == "reopen" for h in clash2.history),
              "回归重开写入流转记录")
        # 上一步待复核的缺洞工单（模型中已消失）-> 自动复核通过
        miss2 = ledger2.get(missing.fingerprint)
        check(miss2.status == STATUS_VERIFIED,
              f"待复核问题消失 -> 自动复核通过（实际 {miss2.status}）")

        # 7b) 完全整改（机电全部绕行/适配）后重新核查 -> 全部闭环、门禁通过
        fixed_dir = os.path.join(td, "fixed")
        make_fixed_sample(fixed_dir, paths)
        result3 = run_coordination(
            [fixed_dir], project="测试项目", owners=owners,
            ledger_path=ledger_path)
        check(result3.summary()["issues_active"] == 0,
              "问题全部消除后未闭环数为 0")
        check(result3.gate_passed, "问题清零后协同门禁通过")
        check(result3.arch_writeback["gate_passed"] is True,
              "回写建筑侧结论为通过")

    # ---- 8) 文件名专业识别 ----
    check(discipline_from_filename("/x/1号楼_结构模型.ifc") == DISC_STRUCT,
          "文件名识别结构专业")
    check(discipline_from_filename("/x/MEP-机电.ifc") == DISC_MEP,
          "文件名识别机电专业")
    check(discipline_from_filename("/x/建筑-A.ifc") == DISC_ARCH,
          "文件名识别建筑专业")

    # ---- 9) 批量核查门禁联动（多专业目录阻断 / 单专业不触发协同）----
    with tempfile.TemporaryDirectory() as td2:
        sample2 = os.path.join(td2, "multi")
        make_coordination_sample(sample2)
        hist = os.path.join(td2, "history")
        batch = run_batch_with_config(
            [sample2], project="批量联动", gate_profile="loose",
            coord_owners={DISC_MEP: "李机电"}, history_dir=hist)
        check(batch.coordination is not None,
              "多专业批量自动执行协同核查")
        check(not batch.gate_passed, "协同未闭环阻断整批放行（即使单体门禁宽松）")
        coord_level = [r for r in batch.gate_results
                       if r.level == "coordination"]
        check(any(not r.passed for r in coord_level),
              "协同门禁规则进入批次放行判定表")

        arch_only = os.path.join(td2, "arch")
        os.makedirs(arch_only)
        arch_src = os.path.join(sample2, "样例-建筑模型.ifc")
        import shutil
        shutil.copy(arch_src, os.path.join(arch_only, "样例-建筑模型.ifc"))
        batch2 = run_batch_with_config(
            [arch_only], project="仅建筑", gate_profile="none",
            history_dir=hist)
        check(batch2.coordination is None,
              "单专业（无机电）不触发协同核查，行为与旧版一致")

    # ---- 10) 参数解析 ----
    s = resolve_settings({"opening_pos_tol_mm": 150})
    check(abs(s.opening_pos_tol - 0.15) < 1e-9, "协同参数毫米 -> 米")
    try:
        resolve_settings({"no_such_key_mm": 1})
        check(False, "非法协同参数应报错")
    except ValueError:
        check(True, "非法协同参数报错")
    check(parse_owner_items(["struct=张工"]) == {"struct": "张工"},
          "责任人参数解析")
    g = resolve_coord_gate("strict", {"coord_max_mismatch_active": 5})
    check(g.max_mismatch_active == 5 and g.require_owner is True,
          "严格协同门禁 + 单项覆盖")
    gs = resolve_coord_gate("default", {"coord_fix_sla_hours": 24})
    check(gs.fix_sla_hours == 24, "整改时限单项覆盖（小时）")
    check(resolve_coord_gate("strict").fix_sla_hours == 48
          and resolve_coord_gate("loose").fix_sla_hours == 168,
          "严格 48h / 宽松 168h 整改时限预设")
    try:
        resolve_coord_gate("default", {"coord_fix_sla_hours": -1})
        check(False, "非法时限应报错")
    except ValueError:
        check(True, "非法整改时限报错")

    print()
    if failures:
        print(f"{len(failures)} 项失败：")
        for m in failures:
            print("  FAIL " + m)
        return 1
    print("全部测试通过。")
    return 0


def make_fixed_sample(fixed_dir, paths):
    """把建筑/结构原样复制到 fixed_dir，机电用“整改后”模型替换。"""
    import shutil
    from tools.make_sample_coordination import _new_file, _pipe, _duct
    os.makedirs(fixed_dir, exist_ok=True)
    shutil.copy(paths["arch"], os.path.join(fixed_dir, "样例-建筑模型.ifc"))
    shutil.copy(paths["struct"], os.path.join(fixed_dir, "样例-结构模型.ifc"))
    f, ctx, st = _new_file("协同样例-机电整改")
    _pipe(f, ctx, st, "P1-给水入户管", 1.5, -0.6, 1.5, 5.5,
          z=2.6, diameter=0.15)
    _duct(f, ctx, st, "P2-排风竖管改路由", 3, 2.0, 0.0, 3, 2.0, 3.0,
          w=0.4, h=0.3)
    _duct(f, ctx, st, "P3-排水小管", 6, 2.5, -0.4, 6, 2.5, 2.6,
          w=0.03, h=0.03)
    _duct(f, ctx, st, "P4-新风管走闲置洞口",
          9.6, -0.5, 2.275, 9.6, 5.0, 2.525, w=0.32, h=0.25)
    f.write(os.path.join(fixed_dir, "样例-机电模型.ifc"))


def make_partial_sample(out_dir, paths):
    """部分整改模型：P4 走闲置洞口（缺洞消失），但 P2 仍穿梁（碰撞保留）。"""
    import shutil
    from tools.make_sample_coordination import _new_file, _pipe, _duct
    os.makedirs(out_dir, exist_ok=True)
    shutil.copy(paths["arch"], os.path.join(out_dir, "样例-建筑模型.ifc"))
    shutil.copy(paths["struct"], os.path.join(out_dir, "样例-结构模型.ifc"))
    f, ctx, st = _new_file("协同样例-机电部分整改")
    _pipe(f, ctx, st, "P1-给水入户管", 1.5, -0.6, 1.5, 5.5,
          z=2.6, diameter=0.15)
    # P2 仍穿梁（原路由 x=3,y=4.5）
    _duct(f, ctx, st, "P2-排风竖管穿梁", 3, 4.5, 0.0, 3, 4.5, 3.0,
          w=0.4, h=0.3)
    # P3 仍为大管（规格不符保留）
    _duct(f, ctx, st, "P3-排水立管", 6, 2.5, -0.4, 6, 2.5, 2.6,
          w=0.3, h=0.3)
    # P4 改走闲置洞口 OP-UNUSED（x=9.6）-> 缺洞消失、闲置洞口被使用
    _duct(f, ctx, st, "P4-新风管走闲置洞口",
          9.6, -0.5, 2.275, 9.6, 5.0, 2.525, w=0.32, h=0.25)
    f.write(os.path.join(out_dir, "样例-机电模型.ifc"))


if __name__ == "__main__":
    raise SystemExit(run())
