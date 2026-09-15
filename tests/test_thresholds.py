"""阈值可配置性测试（不依赖 pytest，可直接运行）。

用法::

    python tests/test_thresholds.py
"""

from __future__ import annotations

import json
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tools.make_sample_ifc import make_sample  # noqa: E402
from ifc_audit.pipeline import audit_ifc_with_config  # noqa: E402
from ifc_audit import report  # noqa: E402
from ifc_audit.thresholds import (  # noqa: E402
    resolve, for_profile, parse_set_items, write_config_template,
    ThresholdConfigError, META,
)
from ifc_audit.cli import main as cli_main  # noqa: E402


def run() -> int:
    failures = []

    def check(cond, msg):
        print(("PASS " if cond else "FAIL ") + msg)
        if not cond:
            failures.append(msg)

    with tempfile.TemporaryDirectory() as td:
        ifc_path = os.path.join(td, "sample.ifc")
        make_sample(ifc_path)

        # ---- 1) 默认预设与样例基线一致 ----
        model = audit_ifc_with_config(ifc_path)
        kinds = {}
        for i in model.issues:
            kinds[i.kind] = kinds.get(i.kind, 0) + 1
        check(kinds.get("wall_end_gap") == 1, "默认预设：1 处墙段缺口")
        check(kinds.get("room_enclosure_gap") == 1, "默认预设：1 处围护缺口")
        check(kinds.get("area_mismatch") == 1, "默认预设：1 个面积偏差")
        check(model.threshold_provenance.is_default(),
              "默认核查的 provenance 标记为默认")
        check(model.thresholds.gap_min_len == 0.10,
              "默认围护缺口下限 100mm")

        # ---- 1b) 回归：门窗归属距离 buffer 不得重复叠加，默认必须 55mm ----
        th0 = model.thresholds
        check(abs(th0.assign_distance - 0.055) < 1e-9,
              f"默认门窗归属距离 = 55mm（实际 {th0.assign_distance*1000:g}mm）")
        check(abs(th0.assign_tol - 0.05) < 1e-9,
              "归属余量 assign_tol = 50mm（barrier_buffer 另算）")
        ts, _ = resolve("strict")
        check(abs(ts.assign_distance - 0.052) < 1e-9,
              f"strict 归属距离 = 52mm（实际 {ts.assign_distance*1000:g}mm）")
        tl, _ = resolve("loose")
        check(abs(tl.assign_distance - 0.060) < 1e-9,
              f"loose 归属距离 = 60mm（实际 {tl.assign_distance*1000:g}mm）")
        # 样例门窗归属数量不受影响
        rooms0 = {r.name: r for r in model.rooms}
        check(rooms0["A-101"].doors == 3 and rooms0["B-102"].doors == 2,
              "样例门归属数量稳定（A=3, B=2）")
        check(rooms0["A-101"].windows == 0 and rooms0["B-102"].windows == 2,
              "样例窗归属数量稳定（B=2，含异常小窗 WN3）")
        # 门窗尺寸下限默认值
        check(model.thresholds.door_min_width == 0.60
              and model.thresholds.door_min_height == 1.80,
              "默认门尺寸下限 600×1800mm")
        check(model.thresholds.win_min_width == 0.40
              and model.thresholds.win_min_height == 0.40,
              "默认窗尺寸下限 400×400mm")

        # ---- 2) 命令行式单项覆盖：聚类容差 100mm → 0.15m 缺口两端不再聚类 ----
        model2 = audit_ifc_with_config(
            ifc_path, overrides=parse_set_items(
                ["endpoint_merge_tol_mm=100"]))
        kinds2 = {}
        for i in model2.issues:
            kinds2[i.kind] = kinds2.get(i.kind, 0) + 1
        check(kinds2.get("wall_end_gap", 0) == 0,
              "聚类容差 100mm：不再判为墙段缺口")
        check(kinds2.get("wall_free_end", 0) == 4,
              f"聚类容差 100mm：缺口两端改报自由端（共 4 个，实际 "
              f"{kinds2.get('wall_free_end', 0)}）")
        check(model2.thresholds.endpoint_merge_tol == 0.10,
              "endpoint_merge_tol 内部换算为 0.10m")
        check(model2.threshold_provenance.profile == "custom",
              "有覆盖时方案标记为 custom")
        check(model2.threshold_provenance.overrides
              == {"endpoint_merge_tol_mm": 100.0},
              "provenance 记录覆盖项")

        # ---- 3) 面积偏差警告线放到 5%：3.75% 偏差不再警告 ----
        model3 = audit_ifc_with_config(
            ifc_path, overrides={"area_dev_warn_pct": 5})
        check(not any(i.kind == "area_mismatch" for i in model3.issues),
              "面积偏差线 5%：样例房间不再警告")
        check(model3.thresholds.area_dev_warn == 0.05,
              "area_dev_warn 内部换算为 0.05")
        mismatch = next(i for i in model.issues if i.kind == "area_mismatch")
        check("警告线 2%" in mismatch.detail,
              "问题详情中注明本次警告线（默认 2%）")

        # ---- 3b) 门窗尺寸下限可配置：放宽到 200mm 后 WN3(300mm) 不再异常 ----
        model3b = audit_ifc_with_config(
            ifc_path, overrides={"win_min_width_mm": 200})
        check(not any(i.kind == "opening_size_anomaly"
                      for i in model3b.issues),
              "窗宽下限 200mm：WN3(300mm) 不再标尺寸异常")
        check(model3b.thresholds.win_min_width == 0.20,
              "win_min_width 内部换算为 0.20m")
        # 收紧到 1300mm：WN1(1200) 与 WN3(300) 都异常，共 2 樘
        model3c = audit_ifc_with_config(
            ifc_path, overrides={"win_min_width_mm": 1300})
        check(sum(1 for i in model3c.issues
                  if i.kind == "opening_size_anomaly") == 2,
              "窗宽下限 1300mm：WN1、WN3 两樘窗标尺寸异常")
        # 未归属与尺寸异常互相独立
        check(sum(1 for i in model3c.issues
                  if i.kind == "opening_unassigned") == 2,
              "放宽/收紧尺寸下限不影响未归属门窗数量（2 樘）")
        # strict 预设下限更严：WN3 仍异常（300 < 500）
        ms = audit_ifc_with_config(ifc_path, profile="strict")
        kinds_s = {}
        for i in ms.issues:
            kinds_s[i.kind] = kinds_s.get(i.kind, 0) + 1
        check(kinds_s.get("opening_size_anomaly") == 1,
              "strict 预设下 WN3 仍为尺寸异常")
        # loose 预设下限更宽：WN3(300) 不再异常
        ml = audit_ifc_with_config(ifc_path, profile="loose")
        check(not any(i.kind == "opening_size_anomaly" for i in ml.issues),
              "loose 预设（窗宽下限 300mm）下无尺寸异常")

        # ---- 4) 配置文件：围护缺口下限 200mm → 150mm 缺口不上报 ----
        cfg_path = os.path.join(td, "th.json")
        with open(cfg_path, "w", encoding="utf-8") as f:
            json.dump({
                "_说明": "测试配置",
                "profile": "default",
                "gap_min_len_mm": 200,
            }, f, ensure_ascii=False)
        model4 = audit_ifc_with_config(ifc_path, config_path=cfg_path)
        check(not any(i.kind == "room_enclosure_gap" for i in model4.issues),
              "配置文件 gap_min_len=200mm：150mm 围护缺口不上报")
        check(model4.threshold_provenance.config_path == cfg_path,
              "provenance 记录配置文件路径")
        gap_issue = next(i for i in model.issues
                         if i.kind == "room_enclosure_gap")
        check("上报下限 100mm" in gap_issue.detail,
              "围护缺口详情注明上报下限")

        # ---- 5) 配置文件指定 strict 预设 ----
        cfg_strict = os.path.join(td, "strict.json")
        with open(cfg_strict, "w", encoding="utf-8") as f:
            json.dump({"profile": "strict"}, f)
        t5, prov5 = resolve(config_path=cfg_strict)
        check(t5.dup_iou == 0.80 and t5.area_dev_warn == 0.01,
              "配置文件 profile=strict 生效（IoU 0.80 / 偏差 1%）")
        check(prov5.profile == "strict", "来源标记为 strict")

        # ---- 6) 预设 + 覆盖的优先级 ----
        t6, prov6 = resolve("strict", overrides={"dup_iou": 0.95})
        check(t6.dup_iou == 0.95 and t6.area_dev_warn == 0.01,
              "--set 在预设基础上覆盖单项，其余保持 strict")
        check(prov6.profile == "custom", "strict 上再覆盖标记为 custom")

        # ---- 7) 库 API 内部名（米/比例）也接受 ----
        t7, _ = resolve(overrides={"free_end_tol": 0.123})
        check(abs(t7.free_end_tol - 0.123) < 1e-9,
              "库 API 传内部名 free_end_tol（米）")

        # ---- 8) 非法配置应报错而不是静默 ----
        for bad in ({"gap_min_len_mm": -1}, {"unknown_key": 1}):
            try:
                resolve(overrides=bad)
                check(False, f"非法阈值 {bad} 应抛 ThresholdConfigError")
            except ThresholdConfigError:
                check(True, f"非法阈值 {bad} 被拒绝")
        try:
            resolve(profile="nope")
            check(False, "未知预设应报错")
        except ThresholdConfigError:
            check(True, "未知预设被拒绝")
        try:
            parse_set_items(["free_end_tol_mm"])
            check(False, "格式错误的 --set 应报错")
        except ThresholdConfigError:
            check(True, "格式错误的 --set 被拒绝")

        # ---- 9) 配置模板生成 / 回读闭环 ----
        tpl_path = os.path.join(td, "template.json")
        write_config_template(tpl_path, "loose")
        t9, prov9 = resolve(config_path=tpl_path)
        check(t9 == for_profile("loose"),
              "loose 模板回读后与 loose 预设完全一致")
        with open(tpl_path, encoding="utf-8") as f:
            tpl = json.load(f)
        check(set(META).issubset(set(tpl)), "模板包含全部阈值键")
        check(all(k.startswith("_") or k == "profile" or k in META
                  for k in tpl), "模板无非注释杂键")
        check(all(k in tpl for k in
                  ("door_min_width_mm", "door_min_height_mm",
                   "win_min_width_mm", "win_min_height_mm")),
              "模板包含门窗尺寸下限四个键")

        # ---- 10) 报告：Excel 含“判定阈值”表与方案行；平面图可导出 ----
        xlsx = report.export_excel(model4, os.path.join(td, "r.xlsx"))
        from openpyxl import load_workbook
        wb = load_workbook(xlsx)
        check("判定阈值" in wb.sheetnames, "Excel 含“判定阈值”工作表")
        check("门窗表" in wb.sheetnames and "门窗明细" in wb.sheetnames,
              "Excel 含“门窗表/门窗明细”工作表")
        ws = wb["判定阈值"]
        first = ws.cell(row=1, column=2).value
        check(cfg_path in (first or ""),
              "阈值表首行注明配置文件来源")
        body = {(r[0], r[1], r[2]) for r in ws.iter_rows(
            min_row=3, values_only=True)}
        check(("房间围护缺口", "围护缺口最小长度", "200 mm") in body,
              "阈值表写入本次实际取值（200 mm）")
        # 门窗表：异常小窗一行 + 未归属两行
        wsop = wb["门窗表"]
        notes_col = [c.value for c in wsop["J"][1:]]
        check(any("尺寸异常" in str(v) for v in notes_col if v),
              "门窗表标注尺寸异常行")
        check(sum("未归属" in str(v) for v in notes_col if v) == 2,
              "门窗表有 2 行未归属门窗")
        summary_vals = [c.value for c in wb["汇总"]["B"]]
        check(any(cfg_path in str(v) for v in summary_vals if v),
              "汇总表含阈值方案行")
        # 门窗表 CSV
        csv_op = report.export_openings_csv(
            model4, os.path.join(td, "op.csv"))
        import csv as _csv
        with open(csv_op, encoding="utf-8-sig") as f:
            rows = list(_csv.DictReader(f))
        check(len(rows) == len(model4.opening_schedule),
              "门窗表 CSV 行数与归并结果一致")
        check({"宽m", "高m", "数量", "备注"}.issubset(rows[0].keys()),
              "门窗表 CSV 含宽/高/数量/备注列")
        png = report.export_annotated_plan(model4,
                                           os.path.join(td, "p.png"))
        check(os.path.getsize(png) > 100, "平面图导出成功")

        # ---- 11) CLI 端到端：--set 与 init-config ----
        out_dir = os.path.join(td, "out")
        rc = cli_main(["audit", ifc_path, "-o", out_dir, "-q",
                       "--set", "area_dev_warn_pct=5"])
        check(rc == 0, "CLI audit --set 退出码 0")
        with open(os.path.join(out_dir, "sample_结果.json"),
                  encoding="utf-8") as f:
            dump = json.load(f)
        check(abs(dump["thresholds"]["values"]["area_dev_warn"] - 0.05)
              < 1e-9,
              "JSON 记录本次阈值（area_dev_warn=0.05）")
        check(dump["thresholds"]["provenance"]["profile"] == "custom",
              "JSON provenance=custom")
        check({"items", "schedule"}.issubset(dump["openings"].keys()),
              "JSON 含 openings 门窗清单段")
        check(len(dump["openings"]["items"]) == 7,
              f"JSON 门窗明细 7 樘（实际 {len(dump['openings']['items'])}）")
        check(any(r["n_unassigned"] for r in dump["openings"]["schedule"]),
              "JSON 门窗表含未归属行")
        check(os.path.exists(os.path.join(out_dir, "sample_门窗表.csv")),
              "CLI 导出门窗表 CSV")

        tpl2 = os.path.join(td, "cli_template.json")
        rc2 = cli_main(["init-config", tpl2, "--profile", "strict"])
        check(rc2 == 0 and os.path.exists(tpl2),
              "CLI init-config 生成 strict 模板")

        rc3 = cli_main(["audit", ifc_path, "-o", out_dir, "-q",
                        "--set", "bad_key=1"])
        check(rc3 == 2, "非法 --set 时退出码 2")

        rc4 = cli_main(["audit", ifc_path, "-o", out_dir, "-q",
                        "--config", os.path.join(td, "missing.json")])
        check(rc4 == 2, "配置文件不存在时退出码 2")

    print()
    if failures:
        print(f"{len(failures)} 项失败")
        return 1
    print("全部测试通过。")
    return 0


if __name__ == "__main__":
    raise SystemExit(run())
